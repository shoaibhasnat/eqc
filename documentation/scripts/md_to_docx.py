"""Convert project Markdown docs to Word (.docx)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]


def set_run_font(run, bold=False, italic=False, code=False, size=11):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Consolas" if code else "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Consolas" if code else "Calibri")
    rFonts.set(qn("w:hAnsi"), "Consolas" if code else "Calibri")
    if code:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_inline_runs(paragraph, text: str, base_size=11):
    # Split `code`, **bold**, *italic* lightly
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, code=True, size=base_size)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True, size=base_size)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, italic=True, size=base_size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size)


def add_paragraph(doc, text: str, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_inline_runs(p, text)
    return p


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        if re.match(r"^\|\s*:?-{3,}", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_inline_runs(p, val, base_size=10)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def md_to_docx(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    in_code = False
    code_lines: list[str] = []
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, parse_table(table_buf))
            table_buf = []

    while i < len(lines):
        line = lines[i]
        raw = line.rstrip("\n")

        if raw.strip().startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, code=True, size=9)
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        if raw.strip().startswith("|"):
            table_buf.append(raw)
            i += 1
            # peek: if next isn't table, flush
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                flush_table()
            continue
        else:
            flush_table()

        if not raw.strip() or raw.strip() == "---":
            i += 1
            continue

        if raw.startswith("# "):
            h = doc.add_heading(raw[2:].strip(), level=0)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0xB8, 0x70, 0x2F)
            i += 1
            continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:].strip(), level=1)
            i += 1
            continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:].strip(), level=2)
            i += 1
            continue
        if raw.startswith("#### "):
            doc.add_heading(raw[5:].strip(), level=3)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", raw.strip())
        if m:
            add_paragraph(doc, m.group(2), style="List Number")
            i += 1
            continue

        if raw.lstrip().startswith("- "):
            add_paragraph(doc, raw.lstrip()[2:].strip(), style="List Bullet")
            i += 1
            continue

        add_paragraph(doc, raw.strip())
        i += 1

    flush_table()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"Wrote {docx_path} ({docx_path.stat().st_size} bytes)")


def main():
    pairs = [
        (
            ROOT / "features" / "app-feature.md",
            ROOT / "features" / "app-feature.docx",
        ),
        (
            ROOT / "complete-project" / "complete-project.md",
            ROOT / "complete-project" / "complete-project.docx",
        ),
    ]
    for md, docx in pairs:
        if not md.exists():
            raise SystemExit(f"Missing {md}")
        md_to_docx(md, docx)


if __name__ == "__main__":
    main()
